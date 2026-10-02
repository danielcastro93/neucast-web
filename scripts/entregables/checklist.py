"""Checklist para el cliente: lo que Neucast necesita de Gabriela y su equipo.

    python3 scripts/entregables/checklist.py <carpeta-de-salida>

Un Excel con el logotipo y los colores de Neucast, para ir marcando qué ya
está. Cada renglón dice qué se necesita, para qué sirve en el sitio y quién
lo resuelve. La columna Estado es una lista: Pendiente, En proceso, Listo.
"""
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
salida = sys.argv[1] if len(sys.argv) > 1 else "."
os.makedirs(salida, exist_ok=True)

TINTA, GRIS, VERDE, SUAVE, LINEA, BLANCO = "1D1D1B", "5F5E59", "48542B", "EDECE7", "E3E1DB", "FFFFFF"

# (grupo, qué necesitamos, para qué sirve, quién)
FILAS = [
    ("Datos de contacto", "Número de WhatsApp de ventas", "Alimenta todos los botones de cotizar del sitio. Hoy está un número de prueba.", "Gabriela"),
    ("Datos de contacto", "Correo al que llegarán las solicitudes del formulario", "El formulario de contacto y la lista de Mi proyecto se mandan a ese correo.", "Gabriela"),
    ("Datos de contacto", "Confirmar las cuentas de Instagram y Facebook", "Van en el pie de página de todo el sitio.", "Gabriela"),
    ("Datos de contacto", "Domicilio, teléfono fijo y horario de atención", "Para la ficha de la empresa en Google. Sin esos tres datos no se puede declarar.", "Gabriela"),
    ("Legales", "Razón social y RFC", "Aviso de privacidad y términos y condiciones.", "Gabriela"),
    ("Legales", "Domicilio fiscal y jurisdicción", "Términos y condiciones.", "Gabriela"),
    ("Legales", "Revisión del aviso de privacidad y los términos por su abogado", "Los textos están en borrador. Describen el sitio real (formulario, Mi proyecto en el navegador, sin cookies ni analítica) y deben validarse antes de publicar.", "Gabriela y su abogado"),
    ("Contenido", "Fotos de producto con fondo transparente (PNG o WebP)", "Las fotos van sobre un fondo gris; con el fondo blanco pegado se vería un cuadro.", "Gabriela"),
    ("Contenido", "Fotos de espacios instalados, por categoría", "Proyectos y bloques de ambiente. Hoy son fotos de muestra.", "Gabriela"),
    ("Contenido", "Fotos de home office con piezas de Neucast", "La sección de home office usa fotos de muestra mientras no haya propias.", "Gabriela"),
    ("Contenido", "Datos de cada pieza según la tabla de campos", "Nombre, tipo, categoría, foto, descripción, material, acabados, disponibilidad y tipo de uso son obligatorios; el resto según el tipo de mueble.", "Gabriela"),
    ("Contenido", "Catálogos propios por categoría, si los tienen (opcional)", "Si pasan un PDF propio, se descarga ese en lugar del generado.", "Gabriela"),
    ("Contenido", "Confirmar por escrito cuántas categorías y piezas van en la carga inicial", "Alcance acordado. La carga del resto del catálogo se cotiza aparte.", "Daniel y Gabriela"),
    ("Decisiones", "¿Entregan a domicilios particulares? ¿A varios domicilios de una misma empresa?", "Define qué promete la sección de home office. Hoy no lo promete.", "Gabriela"),
    ("Decisiones", "¿Pueden enviar muestras de acabados?", "Si sí, se agrega una tarjeta de muestras antes del pie de página.", "Gabriela"),
    ("Decisiones", "¿Tienen modelos 3D o una foto por color de las piezas?", "Fase 2: ver la pieza en el espacio con la cámara y selector de acabados.", "Gabriela"),
    ("Accesos", "Acceso a Hostinger y al registro del dominio neucast.com.mx", "Para publicar el sitio y apuntar el dominio.", "Gabriela"),
    ("Accesos", "Cuenta de Google de la empresa", "Para dar de alta Search Console y, si se decide, analítica.", "Gabriela y Mitch"),
]

wb = Workbook()
ws = wb.active
ws.title = "Checklist"
ws.sheet_view.showGridLines = False

# --- encabezado con logotipo ---
logo = os.path.join(RAIZ, "public", "img", "logo-neucast.png")
if os.path.exists(logo):
    img = Image(logo)
    escala = 44 / img.height
    img.height, img.width = 44, int(img.width * escala)
    ws.add_image(img, "B2")
ws.row_dimensions[2].height = 40
ws["B5"] = "Lo que necesitamos para publicar neucast.com.mx"
ws["B5"].font = Font(name="Helvetica Neue", size=18, bold=True, color=TINTA)
ws["B6"] = f"Checklist para Gabriela · {date.today().strftime('%d/%m/%Y')} · Marca Listo en la columna Estado conforme se vaya resolviendo."
ws["B6"].font = Font(name="Helvetica Neue", size=10, color=GRIS)

# --- tabla ---
cab = ["#", "Grupo", "Qué necesitamos", "Para qué sirve en el sitio", "Quién", "Estado", "Notas"]
anchos = [5, 18, 48, 62, 20, 14, 36]
fila0 = 8
for i, (t, a) in enumerate(zip(cab, anchos), start=2):
    c = ws.cell(row=fila0, column=i, value=t.upper())
    c.font = Font(name="Helvetica Neue", size=8, bold=True, color=BLANCO)
    c.fill = PatternFill("solid", fgColor=TINTA)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.column_dimensions[get_column_letter(i)].width = a
ws.row_dimensions[fila0].height = 24
ws.column_dimensions["A"].width = 3

fina = Side(style="thin", color=LINEA)
grupo_ant = None
for n, (grupo, que, para, quien) in enumerate(FILAS, start=1):
    r = fila0 + n
    valores = [n, grupo if grupo != grupo_ant else "", que, para, quien, "Pendiente", ""]
    for i, v in enumerate(valores, start=2):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name="Helvetica Neue", size=10, color=TINTA if i in (4, 7) else GRIS, bold=(i == 4))
        c.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        c.border = Border(bottom=fina)
        if i == 3 and v:
            c.font = Font(name="Helvetica Neue", size=8, bold=True, color=VERDE)
            c.alignment = Alignment(vertical="top", indent=1)
    ws.row_dimensions[r].height = 44
    grupo_ant = grupo

ultima = fila0 + len(FILAS)
dv = DataValidation(type="list", formula1='"Pendiente,En proceso,Listo"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"G{fila0 + 1}:G{ultima}")

# color por estado (formato condicional con fórmula simple)
from openpyxl.formatting.rule import CellIsRule

ws.conditional_formatting.add(f"G{fila0 + 1}:G{ultima}", CellIsRule(operator="equal", formula=['"Listo"'], fill=PatternFill("solid", fgColor="DCE5C9"), font=Font(color=VERDE, bold=True)))
ws.conditional_formatting.add(f"G{fila0 + 1}:G{ultima}", CellIsRule(operator="equal", formula=['"En proceso"'], fill=PatternFill("solid", fgColor=SUAVE)))

ws.freeze_panes = f"B{fila0 + 1}"
ws.cell(row=ultima + 2, column=2, value="Neucast · neucast.com.mx · Vista previa del sitio: https://danielcastro93.github.io/neucast-web/").font = Font(
    name="Helvetica Neue", size=8, color="95948E"
)
ws.print_options.horizontalCentered = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1

ruta = os.path.join(salida, "neucast-checklist-para-publicar.xlsx")
wb.save(ruta)
print("✓", ruta, f"({len(FILAS)} puntos)")
